#!/usr/bin/env python
"""Fill the report's empty screenshot boxes with the captured images.

Replaces each "FIGURE PLACEHOLDER" table with the real screenshot plus a
caption formatted like the document's existing ones (centred, bold, 13pt).

The concurrency placeholder is removed rather than filled: it calls for a
stress-test capture of the idempotency and Nginx path, which is outside the
machine-learning scope of this report and has no measured result behind it. A
box that cannot be honestly filled should not stay in the document.

Nothing already in the document is modified -- no existing image is touched,
restyled or renumbered.

    python scripts/insert_screenshots.py [--docx vault_api_PBL_report.docx]
"""
from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "report" / "screenshots"
CAPTION_PT = 13

# Matched to the placeholder by a distinctive phrase from its description,
# because table indices shift as soon as one is removed.
FIGURES = [
    (
        "FastAPI Swagger page",
        "fig_swagger_docs.png",
        6.0,
        "Figure 5.1  Interactive API Documentation at /docs",
    ),
    (
        "JSON response body",
        "fig_churn_risk_json.png",
        5.7,
        "Figure 5.2  Churn-Risk Endpoint Response for a High-Risk Merchant",
    ),
    (
        "demonstration dashboard",
        "fig_dashboard_high_risk.png",
        4.4,
        "Figure 5.3  Churn-Risk Dashboard — High-Risk Merchant",
    ),
    (
        "steady recent activity",
        "fig_dashboard_low_risk.png",
        4.4,
        "Figure 5.4  Churn-Risk Dashboard — Low-Risk Merchant",
    ),
]
REMOVE = ["Side-by-side terminal capture"]


def placeholders(document):
    for table in document.tables:
        head = table.rows[0].cells[0].text
        if "FIGURE PLACEHOLDER" in head:
            yield table, head


def find(document, needle: str):
    for table, head in placeholders(document):
        if needle in head:
            return table
    return None


def insert_before(document, table, image: Path, width_in: float, caption: str) -> None:
    """Put a centred image and its caption where the placeholder box was."""
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(image), width=Inches(width_in))
    # Without this the image can land at the foot of one page and leave its
    # caption stranded at the top of the next.
    picture.paragraph_format.keep_with_next = True
    table._element.addprevious(picture._element)

    text = document.add_paragraph()
    text.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = text.add_run(caption)
    run.bold = True
    run.font.size = Pt(CAPTION_PT)
    table._element.addprevious(text._element)

    table._element.getparent().remove(table._element)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--docx", default=str(ROOT / "vault_api_PBL_report.docx"))
    parser.add_argument("--no-backup", action="store_true")
    args = parser.parse_args(argv)

    path = Path(args.docx)
    if not path.exists():
        raise SystemExit(f"not found: {path}")
    missing = [name for _, name, _, _ in FIGURES if not (SHOTS / name).exists()]
    if missing:
        raise SystemExit(
            f"missing screenshots: {missing}. Run scripts/capture_screenshots.py "
            "with the API running."
        )

    if not args.no_backup:
        backup = path.with_suffix(".docx.bak")
        shutil.copy2(path, backup)
        print(f"backup -> {backup.name}", file=sys.stderr)

    document = Document(str(path))
    before = sum(1 for _ in placeholders(document))

    for needle, image, width, caption in FIGURES:
        table = find(document, needle)
        if table is None:
            print(f"  ! no placeholder matching {needle!r}", file=sys.stderr)
            continue
        insert_before(document, table, SHOTS / image, width, caption)
        print(f"  filled  {caption}")

    for needle in REMOVE:
        table = find(document, needle)
        if table is None:
            continue
        table._element.getparent().remove(table._element)
        print(f"  removed placeholder: {needle} (no measured result behind it)")

    document.save(str(path))
    after = sum(1 for _ in placeholders(Document(str(path))))
    print(f"\nplaceholders: {before} -> {after}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
